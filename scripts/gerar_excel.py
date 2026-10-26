#!/usr/bin/env python3
"""
===============================================
Script : gerar_excel.py
Descrição : Consolida todos os JSONs de resultados OCR
            em uma planilha Excel mensal (.xlsx)
Versão : v1.0 - 2026-05-04
Autor  : William Gigliozzi
===============================================

Uso:
    python gerar_excel.py                    # mês atual
    python gerar_excel.py --mes 2026-04      # mês específico
"""

import os
import re
import json
import glob
import argparse
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                  Border, Side, numbers)
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("❌ Execute: pip install openpyxl")

# ── Caminhos (ajuste conforme o ambiente) ─────────────────────
BASE_DIR       = Path(__file__).parent.parent          # raiz do projeto
DIR_RESULTADOS = BASE_DIR / "processados"              # JSONs gerados pelo pipeline
DIR_RELATORIOS = BASE_DIR / "relatorios"               # planilhas de saída
DIR_RELATORIOS.mkdir(exist_ok=True)

# ── Estilos visuais ───────────────────────────────────────────
COR_CABECALHO   = "1F4E79"   # azul escuro
COR_SUBCABEC    = "2E75B6"   # azul médio
COR_LINHA_PAR   = "DEEAF1"   # azul muito claro (linhas pares)
COR_ALERTA      = "FFF2CC"   # amarelo suave (campos não extraídos)
FONTE_BRANCA    = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
FONTE_NORMAL    = Font(name="Calibri", size=10)
FONTE_NEGRITO   = Font(name="Calibri", bold=True, size=10)
BORDA_FINA      = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)


# ── Helpers de estilo ──────────────────────────────────────────

def _estilo_cabecalho(cel, cor=COR_CABECALHO):
    cel.font      = FONTE_BRANCA
    cel.fill      = PatternFill("solid", fgColor=cor)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border    = BORDA_FINA

def _estilo_dado(cel, linha_par=False, alerta=False):
    cel.font      = FONTE_NORMAL
    cor = COR_ALERTA if alerta else (COR_LINHA_PAR if linha_par else "FFFFFF")
    cel.fill      = PatternFill("solid", fgColor=cor)
    cel.alignment = Alignment(vertical="center", wrap_text=True)
    cel.border    = BORDA_FINA

def _ajustar_colunas(ws, larguras: dict):
    for col_letra, largura in larguras.items():
        ws.column_dimensions[col_letra].width = largura


# ── Aba: RENAVAM ───────────────────────────────────────────────

COLUNAS_RENAVAM = [
    ("Arquivo",              18),
    ("Data Consulta",        18),
    ("RENAVAM",              14),
    ("Auto de Infração",     18),
    ("Auto de Renainf",      16),
    ("Data Pagamento c/desc",20),
    ("Data da Infração",     18),
    ("Hora",                  8),
    ("Placa",                10),
    ("Descrição",            40),
    ("Local da Infração",    35),
    ("Valor Original R$",    16),
    ("Valor a Pagar R$",     16),
    ("Status Pagamento",     16),
    ("Tempo OCR (s)",        14),
]

def _criar_aba_renavam(wb: Workbook, registros: list):
    ws = wb.create_sheet("RENAVAM - Nada Consta")
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20

    # Título
    ws.merge_cells("A1:O1")
    cel_titulo = ws["A1"]
    cel_titulo.value     = "RENAVAM — Nada Consta do Veículo"
    cel_titulo.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    cel_titulo.fill      = PatternFill("solid", fgColor=COR_CABECALHO)
    cel_titulo.alignment = Alignment(horizontal="center", vertical="center")

    # Cabeçalhos
    for col, (nome, _) in enumerate(COLUNAS_RENAVAM, 1):
        cel = ws.cell(row=2, column=col, value=nome)
        _estilo_cabecalho(cel, COR_SUBCABEC)

    # Larguras
    for col, (_, larg) in enumerate(COLUNAS_RENAVAM, 1):
        ws.column_dimensions[get_column_letter(col)].width = larg

    # Dados
    linha = 3
    for r in registros:
        for inf in r.get("infracoes", [{}]):
            par = (linha % 2 == 0)
            valores = [
                r.get("arquivo", ""),
                r.get("data_consulta", ""),
                r.get("renavam", ""),
                inf.get("auto_infracao", ""),
                inf.get("auto_renainf", ""),
                inf.get("data_pagamento_desc", ""),
                inf.get("data_infracao", ""),
                inf.get("hora_infracao", ""),
                inf.get("placa_relacionada", ""),
                inf.get("descricao", ""),
                inf.get("local_infracao", ""),
                inf.get("valor_original_rs", ""),
                inf.get("valor_a_pagar_rs", ""),
                inf.get("status_pagamento", ""),
                r.get("tempo_segundos", ""),
            ]
            for col, val in enumerate(valores, 1):
                cel    = ws.cell(row=linha, column=col, value=val)
                alerta = (val == "" and col not in [15])  # tempo é opcional
                _estilo_dado(cel, par, alerta)
                if col in [12, 13]:  # valores monetários
                    cel.number_format = '#.##0,00'
            linha += 1

    # Totalizador simples abaixo dos dados
    if linha > 3:
        ws.cell(row=linha + 1, column=1, value="Total de registros:").font = FONTE_NEGRITO
        ws.cell(row=linha + 1, column=2, value=linha - 3).font = FONTE_NEGRITO

    return ws


# ── Aba: Resumo ────────────────────────────────────────────────

def _criar_aba_resumo(wb: Workbook, registros: list, mes: str):
    ws = wb.create_sheet("Resumo", 0)   # primeira aba
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20

    dados_resumo = [
        ("Mês de referência",     mes),
        ("Total de arquivos",      len(registros)),
        ("Total de infrações",     sum(r.get("total_infracoes", 0) for r in registros)),
        ("Tempo médio OCR (s)",    round(
            sum(r.get("tempo_segundos", 0) for r in registros) / len(registros), 1
        ) if registros else 0),
        ("Gerado em",              datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]

    ws.merge_cells("A1:B1")
    cel = ws["A1"]
    cel.value     = "RESUMO MENSAL — OCR DE DOCUMENTOS"
    cel.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    cel.fill      = PatternFill("solid", fgColor=COR_CABECALHO)
    cel.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for i, (label, valor) in enumerate(dados_resumo, 3):
        cel_l = ws.cell(row=i, column=1, value=label)
        cel_v = ws.cell(row=i, column=2, value=valor)
        cel_l.font   = FONTE_NEGRITO
        cel_v.font   = FONTE_NORMAL
        cel_l.border = BORDA_FINA
        cel_v.border = BORDA_FINA
        cel_l.fill   = PatternFill("solid", fgColor=COR_LINHA_PAR)


# ── Pipeline de geração ────────────────────────────────────────

def gerar(mes: str = None):
    if not mes:
        mes = datetime.now().strftime("%Y-%m")

    print(f"\n📊 Gerando Excel para o mês: {mes}")

    # Carrega todos os JSONs da pasta de processados
    padrao = str(DIR_RESULTADOS / f"*_{mes}*.json")
    arquivos_json = glob.glob(padrao)

    # Se não houver filtro por mês no nome, carrega todos
    if not arquivos_json:
        arquivos_json = glob.glob(str(DIR_RESULTADOS / "*.json"))

    if not arquivos_json:
        print(f"⚠️  Nenhum JSON encontrado em: {DIR_RESULTADOS}")
        return None

    print(f"📂 JSONs encontrados: {len(arquivos_json)}")

    # Separa por tipo de documento
    renavam   = []
    outros    = []

    for path in sorted(arquivos_json):
        with open(path, encoding="utf-8") as f:
            dados = json.load(f)
        tipo = dados.get("tipo_documento", "")
        if "RENAVAM" in tipo.upper():
            renavam.append(dados)
        else:
            outros.append(dados)

    print(f"   RENAVAM : {len(renavam)} arquivo(s)")
    print(f"   Outros  : {len(outros)} arquivo(s)")

    # Cria workbook
    wb = Workbook()
    wb.remove(wb.active)   # remove aba padrão vazia

    todos = renavam + outros
    _criar_aba_resumo(wb, todos, mes)

    if renavam:
        _criar_aba_renavam(wb, renavam)

    # Salva
    nome_arquivo = f"relatorio_ocr_{mes}.xlsx"
    caminho_saida = DIR_RELATORIOS / nome_arquivo
    wb.save(caminho_saida)

    print(f"\n✅ Planilha gerada: {caminho_saida}")
    return str(caminho_saida)


# ── Entrada de linha de comando ────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Gera planilha Excel com resultados OCR")
    parser.add_argument("--mes", default=None,
                        help="Mês no formato AAAA-MM (padrão: mês atual)")
    args = parser.parse_args()
    gerar(args.mes)
