#!/usr/bin/env python3
"""
===============================================
Script : extrator_visao_claude.py
Descrição : Extrai dados de PDFs convertendo
            cada página em imagem e usando
            visão computacional da API Claude.
            Lê texto manuscrito, OCR ruim e
            layouts complexos.
Versão : v1.1 - 2026-10-12
Colunas Excel: Nome | DATA | DOC | COMP | VALOR

⚠️  CONFIGURAÇÃO DA CHAVE DE API:
    Defina a variável de ambiente ANTHROPIC_API_KEY
    antes de executar este script.

    Windows (PowerShell — permanente para o usuário):
        [System.Environment]::SetEnvironmentVariable(
            "ANTHROPIC_API_KEY", "sk-ant-...", "User")

    Windows (sessão atual apenas):
        $env:ANTHROPIC_API_KEY = "sk-ant-..."

    Linux/macOS:
        export ANTHROPIC_API_KEY="sk-ant-..."
===============================================
"""

import base64, io, json, os, re, sys, time, urllib.error, urllib.request
from pathlib import Path
from datetime import datetime

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Garante UTF-8 na saída (no Windows, com pipe, o padrão é cp1252 e quebra em → ✅ ❌)
for _stream in (sys.stdout, sys.stderr):
    if isinstance(_stream, io.TextIOWrapper):
        _stream.reconfigure(encoding="utf-8", errors="replace")

# ── Configuração ──────────────────────────────────────────────
API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
if not API_KEY:
    print("❌ Variável de ambiente ANTHROPIC_API_KEY não definida.")
    print("   Defina a chave antes de executar:")
    print('   PowerShell: $env:ANTHROPIC_API_KEY = "sk-ant-..."')
    print('   Linux/Mac : export ANTHROPIC_API_KEY="sk-ant-..."')
    sys.exit(1)

API_URL    = "https://api.anthropic.com/v1/messages"
MODEL      = "claude-sonnet-4-6"
RESOLUCAO  = 1.5  # fator de zoom da imagem (1.5 = boa qualidade, tamanho razoável)
QUALIDADE  = 85   # qualidade JPEG (0-100)

# ── Estilos Excel ─────────────────────────────────────────────
COR_HEADER = "1F4E79"
COR_PAR    = "DEEAF1"
BORDA = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
COLUNAS = [
    ("Nome",  40),
    ("DATA",  14),
    ("DOC",   22),
    ("COMP",  45),
    ("VALOR", 14),
]

# ── Prompt do sistema ─────────────────────────────────────────
PROMPT_SISTEMA = """Você é um extrator de dados de documentos contábeis brasileiros.
Analise a imagem da página do documento e extraia os campos abaixo. Retorne SOMENTE JSON válido, sem explicações, sem markdown.

Campos de cada registro:
- tipo: classifique o documento (BOLETO, NFSe, NFCe, DAS, DARF, FGTS, HOLERITE, RCI, FATURA_TELEFONE, ENERGIA, AGUA, ISS, SINDICATO, COMPROVANTE_PIX, COMPROVANTE_CARTAO, CHEQUE, EXTRATO_CARTAO, RELATORIO, FERIAS, RESCISAO, AVISO_PREVIO, DESCONHECIDO)
- nome: nome do beneficiário, fornecedor ou empresa emissora — NUNCA inclua CPF ou CNPJ
- data: data de vencimento ou emissão no formato DD/MM/YYYY obrigatoriamente
- doc: número do documento (nota fiscal, boleto, número da fatura, etc.)
- comp: descrição resumida do pagamento — inclua competência, nome do funcionário se houver, serviço ou produto
- valor: valor monetário principal no formato brasileiro (ex: 1.234,56)

ATENÇÃO — MÚLTIPLOS DOCUMENTOS NA MESMA PÁGINA:
- Se a página contiver mais de um documento (ex: vários cheques, vários comprovantes de cartão), retorne um ARRAY JSON com um objeto por documento.
- Se houver apenas um documento, retorne um objeto JSON simples (não array).

Regras importantes:
- Se um campo não existir, retorne string vazia ""
- campo nome: NUNCA inclua CPF (11 dígitos) ou CNPJ (14 dígitos) — extraia apenas o nome
- Para HOLERITE: inclua o nome do funcionário no campo comp (ex: "SALÁRIO SET/2025 - JOAO SILVA")
- Para RCI: inclua o nome do prestador no campo comp (ex: "PRO-LABORE SET/2025 - JOAO SILVA")
- Para COMPROVANTE_CARTAO: nome deve ser o estabelecimento; comp deve indicar bandeira e tipo (ex: "MASTERCARD CRÉDITO À VISTA")
- Para COMPROVANTE_PIX: nome deve ser o destinatário; comp deve indicar o motivo se disponível
- Para CHEQUE: use a data "Bom para" ou "Bom p/" como campo data quando disponível (formato DD/MM/YYYY); se não houver dia use 01 como padrão; gere um registro separado para cada cheque encontrado na página; leia valores escritos à mão
- Para SINDICATO: o campo valor deve conter o valor da contribuição mesmo que apareça como "mensalidade" ou "desconto"
- Para RELATORIO ou NFCe com texto parcial: tente extrair o máximo possível mesmo que incompleto
- datas: SEMPRE no formato DD/MM/YYYY — se encontrar MM/YYYY use 01 como dia
- Retorne SOMENTE o JSON, sem nenhum texto adicional

Exemplos:
- Um documento:   {"tipo":"BOLETO","nome":"SANEAGO","data":"10/11/2025","doc":"123456","comp":"ÁGUA NOV/2025","valor":"187,43"}
- Múltiplos docs: [{"tipo":"CHEQUE","nome":"JOAO SILVA","data":"22/11/2025","doc":"","comp":"CHEQUE BOA PARA NOV/2025","valor":"1.941,32"},{"tipo":"CHEQUE","nome":"JOAO SILVA","data":"22/12/2025","doc":"","comp":"CHEQUE BOA PARA DEZ/2025","valor":"1.941,32"}]"""


# ══════════════════════════════════════════════════════════════
# CONVERSÃO DE PÁGINA PARA IMAGEM BASE64
# ══════════════════════════════════════════════════════════════


def pagina_para_base64(pagina_fitz):
    mat      = fitz.Matrix(RESOLUCAO, RESOLUCAO)
    pix      = pagina_fitz.get_pixmap(matrix=mat)
    img_bytes = pix.tobytes("jpeg", jpg_quality=QUALIDADE)
    return base64.standard_b64encode(img_bytes).decode("utf-8")


# ══════════════════════════════════════════════════════════════
# CHAMADA À API COM IMAGEM
# ══════════════════════════════════════════════════════════════


def chamar_api_visao(img_b64, tentativas=3):
    payload = json.dumps(
        {
            "model": MODEL,
            "max_tokens": 1000,
            "system": PROMPT_SISTEMA,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": "image/jpeg",
                                "data": img_b64,
                            },
                        },
                        {
                            "type": "text",
                            "text": "Extraia os dados desta página conforme as instruções.",
                        },
                    ],
                }
            ],
        }
    ).encode("utf-8")

    for tentativa in range(tentativas):
        try:
            req = urllib.request.Request(
                API_URL,
                data=payload,
                headers={
                    "Content-Type": "application/json",
                    "x-api-key": API_KEY,
                    "anthropic-version": "2023-06-01",
                },
                method="POST",
            )
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = json.loads(resp.read())

            texto  = data["content"][0]["text"].strip()
            texto  = re.sub(r"^```json\s*|\s*```$", "", texto.strip())
            parsed = json.loads(texto)

            return parsed if isinstance(parsed, list) else [parsed]

        except urllib.error.HTTPError as e:
            if e.code in (429, 529):
                wait = 10 * (tentativa + 1)
                print(f" [rate limit, aguardando {wait}s...]", end="", flush=True)
                time.sleep(wait)
            else:
                raise
        except Exception:
            if tentativa == tentativas - 1:
                raise
            time.sleep(3)

    raise Exception("Máximo de tentativas atingido")


# ══════════════════════════════════════════════════════════════
# GERAÇÃO DO EXCEL
# ══════════════════════════════════════════════════════════════


def gerar_excel(registros, caminho_saida):
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title        = "Documentos"
    ws.freeze_panes = "A2"

    for col, (nome, larg) in enumerate(COLUNAS, 1):
        cel           = ws.cell(row=1, column=col, value=nome)
        cel.font      = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cel.fill      = PatternFill("solid", fgColor=COR_HEADER)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border    = BORDA
        ws.column_dimensions[get_column_letter(col)].width = larg
    ws.row_dimensions[1].height = 20

    for i, reg in enumerate(registros, 2):
        par    = i % 2 == 0
        valores = [
            reg.get("nome",  ""),
            reg.get("data",  ""),
            reg.get("doc",   ""),
            reg.get("comp",  ""),
            reg.get("valor", ""),
        ]
        for col, val in enumerate(valores, 1):
            cel           = ws.cell(row=i, column=col, value=val)
            cel.font      = Font(name="Arial", size=10)
            cel.fill      = PatternFill("solid", fgColor=COR_PAR if par else "FFFFFF")
            cel.alignment = Alignment(vertical="center")
            cel.border    = BORDA

    wb.save(caminho_saida)


# ══════════════════════════════════════════════════════════════
# PIPELINE PRINCIPAL
# ══════════════════════════════════════════════════════════════


def processar_pdf(caminho_pdf, caminho_saida=None):
    caminho_pdf = Path(caminho_pdf)
    doc         = fitz.open(str(caminho_pdf))
    total       = len(doc)

    if not caminho_saida:
        data_hora     = datetime.now().strftime("%Y-%m-%d_%H-%M")
        pasta_saida   = Path(__file__).parent.parent / "excel_gerado"
        pasta_saida.mkdir(exist_ok=True)
        caminho_saida = pasta_saida / f"{caminho_pdf.stem}_{data_hora}.xlsx"

    print(f"\n{'='*65}")
    print(f"  Arquivo : {caminho_pdf.name}")
    print(f"  Páginas : {total}")
    print("  Modo    : Visão Computacional (imagem)")
    print(f"  Saída   : {Path(caminho_saida).name}")
    print(f"{'='*65}\n")

    registros      = []
    erros          = []
    contagem_tipos = {}

    for i, pagina in enumerate(doc.pages()):
        num = i + 1
        print(f"  Pág {num:3d}/{total} ", end="", flush=True)

        t0 = time.time()
        try:
            img_b64           = pagina_para_base64(pagina)
            resultados_pagina = chamar_api_visao(img_b64)
            tempo             = round(time.time() - t0, 1)

            for resultado in resultados_pagina:
                tipo = resultado.get("tipo", "DESCONHECIDO")
                nome = resultado.get("nome", "")
                nome = re.sub(r"\s*\d{11}\s*$", "", nome).strip()
                nome = re.sub(r"\s*\d{14}\s*$", "", nome).strip()
                resultado["nome"]   = nome
                resultado["pagina"] = num
                registros.append(resultado)
                contagem_tipos[tipo] = contagem_tipos.get(tipo, 0) + 1

            qtd      = len(resultados_pagina)
            tipo_log = resultados_pagina[0].get("tipo", "?") if resultados_pagina else "?"
            nome_log = (resultados_pagina[0].get("nome", "") or "—")[:35]
            sufixo   = f" (+{qtd-1} doc)" if qtd > 1 else ""
            print(f"→ {tipo_log:<22} {nome_log:<36} {tempo}s{sufixo}")

        except Exception as e:
            tempo = round(time.time() - t0, 1)
            print(f"→ ERRO: {str(e)[:50]} {tempo}s")
            erros.append({"pagina": num, "erro": str(e)})

        time.sleep(0.3)

    doc.close()

    vistos           = set()
    registros_unicos = []
    for r in registros:
        chave = (r.get("tipo"), r.get("valor"), r.get("data"), r.get("nome"))
        if chave not in vistos:
            vistos.add(chave)
            registros_unicos.append(r)

    gerar_excel(registros_unicos, str(caminho_saida))

    print(f"\n{'='*65}")
    print("  RESUMO")
    print(f"{'='*65}")
    for tipo, qtd in sorted(contagem_tipos.items(), key=lambda x: -x[1]):
        print(f"  {tipo:<32}: {qtd}")
    print(f"{'─'*65}")
    print(f"  Total registros  : {len(registros_unicos)}")
    if erros:
        print(f"  Erros            : {len(erros)} páginas")
        for e in erros:
            print(f"    Pág {e['pagina']}: {e['erro']}")
    print(f"{'='*65}")
    print(f"\n✅ Excel salvo em: {caminho_saida}\n")

    pasta_logs = Path(__file__).parent.parent / "logs"
    pasta_logs.mkdir(exist_ok=True)
    json_saida = pasta_logs / f"{Path(caminho_saida).stem}.json"
    with open(json_saida, "w", encoding="utf-8") as f:
        json.dump(registros_unicos, f, ensure_ascii=False, indent=2)
    print(f"📋 JSON de auditoria: {json_saida}\n")

    return registros_unicos


# ── Linha de comando ──────────────────────────────────────────


def main():
    if len(sys.argv) < 2:
        print("Uso: python extrator_visao_claude.py <arquivo.pdf> [saida.xlsx]")
        print("     python extrator_visao_claude.py pasta/entrada/")
        sys.exit(0)

    entrada = Path(sys.argv[1])
    saida   = Path(sys.argv[2]) if len(sys.argv) >= 3 else None

    if entrada.is_dir():
        pdfs = sorted(entrada.glob("*.pdf")) + sorted(entrada.glob("*.PDF"))
        if not pdfs:
            print(f"Nenhum PDF encontrado em: {entrada}")
            sys.exit(1)
        for pdf in pdfs:
            processar_pdf(pdf)
    elif entrada.is_file():
        processar_pdf(entrada, saida)
    else:
        print(f"❌ Caminho não encontrado: {entrada}")
        sys.exit(1)


if __name__ == "__main__":
    main()
